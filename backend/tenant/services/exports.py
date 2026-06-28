from __future__ import annotations

import csv
import json
import os
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
from typing import Any, Iterable

from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tenant.models import ExportJob


SUPPORTED_EXPORT_FORMATS = {"csv", "xlsx", "json", "pdf"}
EXPORT_SYNC_ROW_LIMIT = int(getattr(settings, "EXPORT_SYNC_ROW_LIMIT", 1000))

CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "pdf": "application/pdf",
}

EXTENSIONS = {
    "csv": "csv",
    "xlsx": "xlsx",
    "json": "json",
    "pdf": "pdf",
}


@dataclass(frozen=True)
class ExportDataset:
    title: str
    report_type: str
    columns: list[dict[str, str]]
    rows: Iterable[dict[str, Any]]
    filename_prefix: str
    summary: list[dict[str, Any]] | None = None
    applied_filters: dict[str, Any] | None = None
    window: dict[str, Any] | None = None
    row_count: int | None = None


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    filename: str
    content_type: str


def normalize_export_format(export_format: str | None, *, allow_pdf: bool = True) -> str:
    normalized = (export_format or "csv").strip().lower()
    allowed = SUPPORTED_EXPORT_FORMATS if allow_pdf else SUPPORTED_EXPORT_FORMATS - {"pdf"}
    if normalized not in allowed:
        choices = ", ".join(sorted(allowed))
        raise ValueError(f"Unsupported export format. Use one of: {choices}.")
    return normalized


def should_queue_export(row_count: int) -> bool:
    return row_count > EXPORT_SYNC_ROW_LIMIT


def is_admin_user(user) -> bool:
    role_name = (getattr(getattr(user, "role", None), "name", "") or "").strip().lower()
    return bool(getattr(user, "is_authenticated", False) and role_name == "admin")


def build_export_file(dataset: ExportDataset, export_format: str) -> ExportFile:
    export_format = normalize_export_format(export_format)
    if export_format == "csv":
        content = _csv_bytes(dataset.columns, dataset.rows)
    elif export_format == "xlsx":
        content = _xlsx_bytes(dataset)
    elif export_format == "json":
        content = _json_bytes(dataset)
    else:
        content = _pdf_bytes(dataset)

    filename = f"{dataset.filename_prefix}.{EXTENSIONS[export_format]}"
    return ExportFile(content=content, filename=filename, content_type=CONTENT_TYPES[export_format])


def export_file_response(export_file: ExportFile) -> HttpResponse:
    response = HttpResponse(export_file.content, content_type=export_file.content_type)
    response["Content-Disposition"] = f'attachment; filename="{export_file.filename}"'
    return response


def queue_export_job(
    *,
    user,
    export_type: str,
    export_format: str,
    params: dict[str, Any] | None = None,
    selected_ids: list[int] | None = None,
    row_count: int = 0,
    request=None,
) -> ExportJob:
    from tenant.tasks import generate_export_job
    from tenant.services.audit import actor_name, log_audit_event

    job = ExportJob.objects.create(
        business=user.business,
        requested_by=user,
        export_type=export_type,
        format=export_format,
        params=params or {},
        selected_ids=selected_ids or [],
        row_count=row_count,
    )
    log_audit_event(
        request,
        event_type="REPORT_EXPORT_QUEUED",
        entity_type="report",
        entity=job,
        activity=f"{actor_name(user)} queued {export_type} export",
        actor_user=user,
        business=user.business,
        metadata=_audit_metadata(export_type=export_type, export_format=export_format, row_count=row_count, params=params),
    )
    generate_export_job.delay(job.id)
    return job


def build_export_dataset_for_job(job: ExportJob, *, stream: bool = False) -> ExportDataset:
    from tenant.services.reports import (
        build_report_export_dataset,
        build_selected_task_export_dataset,
        build_selected_ticket_export_dataset,
    )

    if job.export_type in {"tickets", "tasks", "agents"}:
        return build_report_export_dataset(
            job.requested_by,
            job.export_type,
            job.params,
            stream=stream,
            business=job.business,
        )
    elif job.export_type == "selected_tickets":
        return build_selected_ticket_export_dataset(
            job.requested_by,
            job.selected_ids,
            stream=stream,
            business=job.business,
        )
    elif job.export_type == "selected_tasks":
        return build_selected_task_export_dataset(
            job.requested_by,
            job.selected_ids,
            stream=stream,
            business=job.business,
        )
    raise ValueError(f"Unsupported export job type: {job.export_type}")


def generate_export_file_for_job(job: ExportJob) -> ExportFile:
    return build_export_file(build_export_dataset_for_job(job), job.format)


def save_export_file_for_job(job: ExportJob) -> ExportFile:
    dataset = build_export_dataset_for_job(job, stream=True)
    filename = f"{dataset.filename_prefix}.{EXTENSIONS[job.format]}"
    export_file = ExportFile(content=b"", filename=filename, content_type=CONTENT_TYPES[job.format])
    file_path = _export_file_path(job, filename)
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    _write_export_file_to_path(dataset, job.format, file_path)
    job.file_path = file_path
    return export_file


def save_export_file(job: ExportJob, export_file: ExportFile) -> str:
    file_path = _export_file_path(job, export_file.filename)
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    with open(file_path, "wb") as output_file:
        output_file.write(export_file.content)
    return file_path


def audit_export_completed(job: ExportJob) -> None:
    from tenant.services.audit import actor_name, log_audit_event

    log_audit_event(
        None,
        event_type="REPORT_EXPORT_COMPLETED",
        entity_type="report",
        entity=job,
        activity=f"{actor_name(job.requested_by)} completed {job.export_type} export",
        actor_user=job.requested_by,
        business=job.business,
        metadata=_audit_metadata(
            export_type=job.export_type,
            export_format=job.format,
            row_count=job.row_count,
            filename=job.filename,
        ),
    )


def audit_export_failed(job: ExportJob, error: str) -> None:
    from tenant.services.audit import actor_name, log_audit_event

    log_audit_event(
        None,
        event_type="REPORT_EXPORT_FAILED",
        entity_type="report",
        entity=job,
        activity=f"{actor_name(job.requested_by)} failed {job.export_type} export",
        actor_user=job.requested_by,
        business=job.business,
        metadata=_audit_metadata(
            export_type=job.export_type,
            export_format=job.format,
            row_count=job.row_count,
            error=error,
        ),
    )


def audit_export_downloaded(request, job: ExportJob) -> None:
    from tenant.services.audit import actor_name, log_audit_event

    log_audit_event(
        request,
        event_type="REPORT_EXPORT_DOWNLOADED",
        entity_type="report",
        entity=job,
        activity=f"{actor_name(request.user)} downloaded {job.export_type} export",
        metadata=_audit_metadata(
            export_type=job.export_type,
            export_format=job.format,
            row_count=job.row_count,
            filename=job.filename,
        ),
    )


def audit_direct_export(request, *, export_type: str, export_format: str, row_count: int, params: dict[str, Any] | None = None) -> None:
    from tenant.services.audit import actor_name, log_audit_event

    log_audit_event(
        request,
        event_type="REPORT_EXPORTED",
        entity_type="report",
        entity=request.user.business,
        activity=f"{actor_name(request.user)} exported {export_type} report",
        metadata=_audit_metadata(export_type=export_type, export_format=export_format, row_count=row_count, params=params),
    )


def _audit_metadata(
    *,
    export_type: str,
    export_format: str,
    row_count: int,
    params: dict[str, Any] | None = None,
    filename: str | None = None,
    error: str | None = None,
) -> dict[str, Any]:
    metadata: dict[str, Any] = {
        "report_type": export_type,
        "format": export_format,
        "row_count": row_count,
    }
    if filename:
        metadata["filename"] = filename
    if error:
        metadata["error"] = error[:500]
    if params:
        metadata["filters"] = {
            key: value
            for key, value in params.items()
            if key not in {"page", "page_size", "format"} and value not in [None, "", "all"]
        }
    return metadata


def _export_file_path(job: ExportJob, filename: str) -> str:
    return os.path.join(settings.MEDIA_ROOT, "exports", str(job.business_id), f"{job.id}-{filename}")


def export_job_payload(job: ExportJob) -> dict[str, Any]:
    return {
        "id": job.id,
        "export_id": job.id,
        "status": job.job_status,
        "export_type": job.export_type,
        "format": job.format,
        "row_count": job.row_count,
        "filename": job.filename,
        "content_type": job.content_type,
        "error": job.error,
        "created_at": timezone.localtime(job.created_at).isoformat() if job.created_at else None,
        "completed_at": timezone.localtime(job.completed_at).isoformat() if job.completed_at else None,
    }


def _csv_bytes(columns: list[dict[str, str]], rows: list[dict[str, Any]]) -> bytes:
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([column["label"] for column in columns])
    for row in rows:
        writer.writerow([_cell_value(row.get(column["key"], "")) for column in columns])
    return output.getvalue().encode("utf-8-sig")


def _json_bytes(dataset: ExportDataset) -> bytes:
    payload = {
        "title": dataset.title,
        "report_type": dataset.report_type,
        "generated_at": timezone.localtime(timezone.now()).isoformat(),
        "window": dataset.window or {},
        "applied_filters": dataset.applied_filters or {},
        "summary": dataset.summary or [],
        "columns": dataset.columns,
        "rows": dataset.rows,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def _xlsx_bytes(dataset: ExportDataset) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    data_sheet = workbook.create_sheet("Data")
    filters_sheet = workbook.create_sheet("Filters")

    _write_summary_sheet(summary_sheet, dataset)
    _write_filters_sheet(filters_sheet, dataset.applied_filters or {}, dataset.window or {})
    _write_data_sheet(data_sheet, dataset.columns, dataset.rows)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_export_file_to_path(dataset: ExportDataset, export_format: str, file_path: str) -> None:
    if export_format == "csv":
        _write_csv_file(dataset.columns, dataset.rows, file_path)
    elif export_format == "json":
        _write_json_file(dataset, file_path)
    elif export_format == "xlsx":
        _write_xlsx_file(dataset, file_path)
    elif export_format == "pdf":
        file_bytes = _pdf_bytes(dataset)
        with open(file_path, "wb") as output_file:
            output_file.write(file_bytes)
    else:
        raise ValueError(f"Unsupported export format: {export_format}")


def _write_csv_file(columns: list[dict[str, str]], rows: Iterable[dict[str, Any]], file_path: str) -> None:
    with open(file_path, "w", newline="", encoding="utf-8-sig") as output_file:
        writer = csv.writer(output_file)
        writer.writerow([column["label"] for column in columns])
        for row in rows:
            writer.writerow([_cell_value(row.get(column["key"], "")) for column in columns])


def _write_json_file(dataset: ExportDataset, file_path: str) -> None:
    metadata = {
        "title": dataset.title,
        "report_type": dataset.report_type,
        "generated_at": timezone.localtime(timezone.now()).isoformat(),
        "window": dataset.window or {},
        "applied_filters": dataset.applied_filters or {},
        "summary": dataset.summary or [],
        "columns": dataset.columns,
    }
    with open(file_path, "w", encoding="utf-8") as output_file:
        output_file.write("{\n")
        for index, (key, value) in enumerate(metadata.items()):
            output_file.write(f'  "{key}": {json.dumps(value, ensure_ascii=False, default=str)},\n')
        output_file.write('  "rows": [\n')
        first = True
        for row in dataset.rows:
            if not first:
                output_file.write(",\n")
            output_file.write(f"    {json.dumps(row, ensure_ascii=False, default=str)}")
            first = False
        output_file.write("\n  ]\n}\n")


def _write_xlsx_file(dataset: ExportDataset, file_path: str) -> None:
    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet("Summary")
    filters_sheet = workbook.create_sheet("Filters")
    data_sheet = workbook.create_sheet("Data")

    _append_write_only_row(summary_sheet, [dataset.title])
    _append_write_only_row(summary_sheet, ["Generated At", timezone.localtime(timezone.now()).strftime("%Y-%m-%d %H:%M")])
    _append_write_only_row(summary_sheet, [])
    _append_write_only_row(summary_sheet, ["Metric", "Value"], header=True)
    for metric in dataset.summary or []:
        _append_write_only_row(summary_sheet, [metric.get("label"), _cell_value(metric.get("display_value"))])

    _append_write_only_row(filters_sheet, ["Applied Filters"])
    _append_write_only_row(filters_sheet, [])
    _append_write_only_row(filters_sheet, ["Field", "Value"], header=True)
    for key, value in {**(dataset.window or {}), **(dataset.applied_filters or {})}.items():
        _append_write_only_row(filters_sheet, [str(key).replace("_", " ").title(), _cell_value(value)])

    _append_write_only_row(data_sheet, [column["label"] for column in dataset.columns], header=True)
    for row in dataset.rows:
        _append_write_only_row(data_sheet, [_cell_value(row.get(column["key"], "")) for column in dataset.columns])

    workbook.save(file_path)


def _append_write_only_row(sheet, values: list[Any], *, header: bool = False) -> None:
    if not header:
        sheet.append(values)
        return
    fill = PatternFill("solid", fgColor="E5E7EB")
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(bold=True, color="111827")
        cell.fill = fill
        cells.append(cell)
    sheet.append(cells)


def _write_summary_sheet(sheet, dataset: ExportDataset) -> None:
    sheet["A1"] = dataset.title
    sheet["A1"].font = Font(bold=True, size=16, color="1F2937")
    sheet["A2"] = "Generated At"
    sheet["B2"] = timezone.localtime(timezone.now()).strftime("%Y-%m-%d %H:%M")
    sheet["A4"] = "Metric"
    sheet["B4"] = "Value"
    _style_header_row(sheet, 4, 2)

    for row_index, metric in enumerate(dataset.summary or [], start=5):
        sheet.cell(row=row_index, column=1, value=metric.get("label"))
        sheet.cell(row=row_index, column=2, value=_cell_value(metric.get("display_value")))

    _size_columns(sheet, max_row=max(5, len(dataset.summary or []) + 5), max_col=2)


def _write_filters_sheet(sheet, applied_filters: dict[str, Any], window: dict[str, Any]) -> None:
    sheet["A1"] = "Applied Filters"
    sheet["A1"].font = Font(bold=True, size=14, color="1F2937")
    sheet["A3"] = "Field"
    sheet["B3"] = "Value"
    _style_header_row(sheet, 3, 2)

    rows = {**window, **applied_filters}
    for row_index, (key, value) in enumerate(rows.items(), start=4):
        sheet.cell(row=row_index, column=1, value=str(key).replace("_", " ").title())
        sheet.cell(row=row_index, column=2, value=_cell_value(value))

    _size_columns(sheet, max_row=max(4, len(rows) + 4), max_col=2)


def _write_data_sheet(sheet, columns: list[dict[str, str]], rows: list[dict[str, Any]]) -> None:
    for column_index, column in enumerate(columns, start=1):
        sheet.cell(row=1, column=column_index, value=column["label"])
    _style_header_row(sheet, 1, len(columns))

    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(columns, start=1):
            sheet.cell(row=row_index, column=column_index, value=_cell_value(row.get(column["key"], "")))

    if columns:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"
    _size_columns(sheet, max_row=len(rows) + 1, max_col=len(columns))


def _style_header_row(sheet, row_index: int, column_count: int) -> None:
    fill = PatternFill("solid", fgColor="E5E7EB")
    for column_index in range(1, column_count + 1):
        cell = sheet.cell(row=row_index, column=column_index)
        cell.font = Font(bold=True, color="111827")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _size_columns(sheet, *, max_row: int, max_col: int) -> None:
    for column_index in range(1, max_col + 1):
        width = 12
        for row_index in range(1, max_row + 1):
            value = sheet.cell(row=row_index, column=column_index).value
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 48))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _pdf_bytes(dataset: ExportDataset) -> bytes:
    lines: list[str] = [
        dataset.title,
        f"Date Window: {(dataset.window or {}).get('label', 'Export')}",
        "",
        "Summary",
    ]
    for metric in dataset.summary or []:
        lines.append(f"- {metric.get('label')}: {metric.get('display_value')}")

    lines.extend(["", "Applied Filters"])
    for key, value in (dataset.applied_filters or {}).items():
        if value in (None, "", "all"):
            continue
        lines.append(f"- {str(key).replace('_', ' ').title()}: {value}")

    lines.extend(["", "Rows"])
    for index, row in enumerate(dataset.rows, start=1):
        lines.append(f"{index}. {_primary_row_title(dataset.report_type, row)}")
        for column in dataset.columns:
            lines.append(f"   {column['label']}: {_cell_value(row.get(column['key'], '-'))}")
        lines.append("")

    return _build_simple_pdf(lines)


def _build_simple_pdf(lines: list[str]) -> bytes:
    max_lines_per_page = 48
    pages = [lines[index:index + max_lines_per_page] for index in range(0, len(lines), max_lines_per_page)] or [[]]
    font_id = 1
    first_content_id = 2
    first_page_id = first_content_id + len(pages)
    pages_id = first_page_id + len(pages)
    catalog_id = pages_id + 1
    objects: dict[int, bytes] = {font_id: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}

    for index, page_lines in enumerate(pages):
        content_id = first_content_id + index
        page_id = first_page_id + index
        stream = _pdf_content_stream(page_lines)
        objects[content_id] = f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream"
        objects[page_id] = (
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode("latin-1")

    kids = " ".join(f"{first_page_id + index} 0 R" for index in range(len(pages)))
    objects[pages_id] = f"<< /Type /Pages /Kids [ {kids} ] /Count {len(pages)} >>".encode("latin-1")
    objects[catalog_id] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("latin-1")

    payload = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0] * (catalog_id + 1)
    for object_id in range(1, catalog_id + 1):
        offsets[object_id] = len(payload)
        payload.extend(f"{object_id} 0 obj\n".encode("latin-1"))
        payload.extend(objects[object_id])
        payload.extend(b"\nendobj\n")

    xref_offset = len(payload)
    payload.extend(f"xref\n0 {catalog_id + 1}\n".encode("latin-1"))
    payload.extend(b"0000000000 65535 f \n")
    for object_id in range(1, catalog_id + 1):
        payload.extend(f"{offsets[object_id]:010d} 00000 n \n".encode("latin-1"))
    payload.extend(
        (
            f"trailer\n<< /Size {catalog_id + 1} /Root {catalog_id} 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF"
        ).encode("latin-1")
    )
    return bytes(payload)


def _pdf_content_stream(lines: list[str]) -> bytes:
    chunks = ["BT", "/F1 9 Tf", "40 752 Td"]
    for index, line in enumerate(lines):
        escaped = _escape_pdf_text(line[:180])
        if index == 0:
            chunks.append(f"({escaped}) Tj")
        else:
            chunks.append("0 -14 Td")
            chunks.append(f"({escaped}) Tj")
    chunks.append("ET")
    return "\n".join(chunks).encode("latin-1", errors="replace")


def _escape_pdf_text(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _primary_row_title(report_type: str, row: dict[str, Any]) -> str:
    if report_type == "tickets":
        return str(row.get("ticket_id", "Ticket"))
    if report_type == "tasks":
        return str(row.get("task_id", "Task"))
    return str(row.get("agent_name", "Agent"))


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")
    return value
